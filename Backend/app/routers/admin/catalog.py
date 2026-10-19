"""Admin: categories, brands, products, variants, images, bulk import/export (AD-02 … AD-05)."""
import csv
import io
import re
from decimal import Decimal, InvalidOperation

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, selectinload

from app.core.deps import staff_or_admin
from app.db import get_db
from app.models import Brand, Category, OrderItem, Product, ProductImage, ProductVariant, User
from app.schemas import GENDERS, BrandIn, CategoryIn, ImageOrderIn, ProductIn, ReorderIn, check_gst_rate
from app.serializers import brand_out, category_out, product_card, product_detail
from app.services.audit import audit
from app.services.media import delete_media, save_image
from app.services.stock import change_stock
from app.utils import unique_slug

router = APIRouter(dependencies=[Depends(staff_or_admin)])

MAX_IMAGES = 8  # per product; at least one is mandatory


# ------------------------------------------------------------------ generic upload

@router.post("/upload")
async def upload(folder: str = Query(pattern="^(categories|brands|banners|store)$"), file: UploadFile = File(...)):
    return {"url": await save_image(file, folder)}


# ------------------------------------------------------------------ categories

@router.get("/categories")
def list_categories(db: Session = Depends(get_db)):
    rows = db.scalars(select(Category).order_by(Category.sort_order, Category.name)).all()
    counts = dict(db.execute(select(Product.category_id, func.count(Product.id)).group_by(Product.category_id)).all())
    return [{**category_out(c), "product_count": counts.get(c.id, 0)} for c in rows]


@router.post("/categories", status_code=201)
def create_category(body: CategoryIn, request: Request, db: Session = Depends(get_db), user: User = Depends(staff_or_admin)):
    if body.parent_id and not db.get(Category, body.parent_id):
        raise HTTPException(400, "Parent category not found")
    c = Category(**body.model_dump(exclude={"slug"}), slug=unique_slug(db, Category, body.slug or body.name))
    db.add(c)
    db.flush()
    audit(db, user, "create", "category", c.id, {"name": c.name}, request)
    db.commit()
    return category_out(c)


@router.put("/categories/{cid}")
def update_category(cid: int, body: CategoryIn, request: Request, db: Session = Depends(get_db), user: User = Depends(staff_or_admin)):
    c = db.get(Category, cid)
    if not c:
        raise HTTPException(404, "Category not found")
    if body.parent_id == cid:
        raise HTTPException(400, "A category cannot be its own parent")
    for k, v in body.model_dump(exclude={"slug"}).items():
        setattr(c, k, v)
    c.slug = unique_slug(db, Category, body.slug or body.name, exclude_id=c.id)
    audit(db, user, "update", "category", c.id, body.model_dump(), request)
    db.commit()
    return category_out(c)


@router.delete("/categories/{cid}")
def delete_category(cid: int, request: Request, db: Session = Depends(get_db), user: User = Depends(staff_or_admin)):
    c = db.get(Category, cid)
    if not c:
        raise HTTPException(404, "Category not found")
    if db.scalar(select(func.count(Product.id)).where(Product.category_id == cid)):
        raise HTTPException(400, "Move or delete the products in this category first")
    for child in db.scalars(select(Category).where(Category.parent_id == cid)).all():
        child.parent_id = c.parent_id
    db.delete(c)
    audit(db, user, "delete", "category", cid, {"name": c.name}, request)
    db.commit()
    return {"ok": True}


@router.post("/categories/reorder")
def reorder_categories(body: ReorderIn, db: Session = Depends(get_db)):
    for i, cid in enumerate(body.ids):
        c = db.get(Category, cid)
        if c:
            c.sort_order = i
    db.commit()
    return {"ok": True}


# ------------------------------------------------------------------ brands

@router.get("/brands")
def list_brands(db: Session = Depends(get_db)):
    rows = db.scalars(select(Brand).order_by(Brand.name)).all()
    counts = dict(db.execute(select(Product.brand_id, func.count(Product.id)).group_by(Product.brand_id)).all())
    return [{**brand_out(b), "product_count": counts.get(b.id, 0)} for b in rows]


@router.post("/brands", status_code=201)
def create_brand(body: BrandIn, request: Request, db: Session = Depends(get_db), user: User = Depends(staff_or_admin)):
    b = Brand(**body.model_dump(exclude={"slug"}), slug=unique_slug(db, Brand, body.slug or body.name))
    db.add(b)
    db.flush()
    audit(db, user, "create", "brand", b.id, {"name": b.name}, request)
    db.commit()
    return brand_out(b)


@router.put("/brands/{bid}")
def update_brand(bid: int, body: BrandIn, request: Request, db: Session = Depends(get_db), user: User = Depends(staff_or_admin)):
    b = db.get(Brand, bid)
    if not b:
        raise HTTPException(404, "Brand not found")
    for k, v in body.model_dump(exclude={"slug"}).items():
        setattr(b, k, v)
    b.slug = unique_slug(db, Brand, body.slug or body.name, exclude_id=b.id)
    audit(db, user, "update", "brand", b.id, body.model_dump(), request)
    db.commit()
    return brand_out(b)


@router.delete("/brands/{bid}")
def delete_brand(bid: int, request: Request, db: Session = Depends(get_db), user: User = Depends(staff_or_admin)):
    b = db.get(Brand, bid)
    if not b:
        raise HTTPException(404, "Brand not found")
    if db.scalar(select(func.count(Product.id)).where(Product.brand_id == bid)):
        raise HTTPException(400, "This brand has products. Deactivate it instead.")
    db.delete(b)
    audit(db, user, "delete", "brand", bid, {"name": b.name}, request)
    db.commit()
    return {"ok": True}


# ------------------------------------------------------------------ products

def _load(db: Session, pid: int) -> Product:
    p = db.scalar(
        select(Product).options(
            selectinload(Product.images), selectinload(Product.variants),
            selectinload(Product.brand), selectinload(Product.category),
        ).where(Product.id == pid)
    )
    if not p:
        raise HTTPException(404, "Product not found")
    return p


@router.get("/products")
def list_products(
    db: Session = Depends(get_db),
    q: str | None = None,
    category_id: int | None = None,
    brand_id: int | None = None,
    status: str | None = Query(default=None, pattern="^(active|draft)$"),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=25, ge=1, le=100),
):
    stmt = select(Product)
    if q:
        like = f"%{q.strip()}%"
        stmt = stmt.where(or_(
            Product.name.like(like), Product.slug.like(like),
            Product.id.in_(select(ProductVariant.product_id).where(ProductVariant.sku.like(like))),
        ))
    if category_id:
        stmt = stmt.where(Product.category_id == category_id)
    if brand_id:
        stmt = stmt.where(Product.brand_id == brand_id)
    if status:
        stmt = stmt.where(Product.is_active.is_(status == "active"))
    total = db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    rows = db.scalars(
        stmt.options(selectinload(Product.images), selectinload(Product.variants), selectinload(Product.brand), selectinload(Product.category))
        .order_by(Product.id.desc()).offset((page - 1) * page_size).limit(page_size)
    ).all()
    items = []
    for p in rows:
        card = product_card(p)
        card.update(is_active=p.is_active, total_stock=sum(v.stock_qty for v in p.variants), variant_count=len(p.variants))
        items.append(card)
    return {"items": items, "total": total, "page": page, "page_size": page_size}


@router.get("/products/{pid}")
def get_product(pid: int, db: Session = Depends(get_db)):
    return product_detail(_load(db, pid), admin=True)


def _check_refs(db: Session, body: ProductIn) -> None:
    if body.brand_id and not db.get(Brand, body.brand_id):
        raise HTTPException(400, "Brand not found")
    if body.category_id and not db.get(Category, body.category_id):
        raise HTTPException(400, "Category not found")
    keys = [(v.size.strip().lower(), v.colour.strip().lower()) for v in body.variants]
    if len(keys) != len(set(keys)):
        raise HTTPException(400, "Each size + colour combination must be unique")
    skus = [v.sku.strip().upper() for v in body.variants]
    if len(skus) != len(set(skus)):
        raise HTTPException(400, "SKUs must be unique")


def _sku_taken(db: Session, sku: str, exclude_variant_id: int | None) -> bool:
    stmt = select(ProductVariant.id).where(func.upper(ProductVariant.sku) == sku.upper())
    if exclude_variant_id:
        stmt = stmt.where(ProductVariant.id != exclude_variant_id)
    return bool(db.scalar(stmt))


@router.post("/products", status_code=201)
async def create_product(
    request: Request,
    data: str = Form(..., description="The product as ProductIn JSON"),
    files: list[UploadFile] | None = File(default=None, description=f"1 to {MAX_IMAGES} product images (JPG, PNG or WebP)"),
    db: Session = Depends(get_db),
    user: User = Depends(staff_or_admin),
):
    """Creates the product, its variants and its images in one step; at least one image is mandatory."""
    try:
        body = ProductIn.model_validate_json(data)
    except ValidationError as exc:
        raise RequestValidationError(exc.errors(include_url=False, include_context=False)) from exc
    files = [f for f in files or [] if f.filename]
    if not files:
        raise HTTPException(422, "Add at least one product image")
    if len(files) > MAX_IMAGES:
        raise HTTPException(400, f"A product can have at most {MAX_IMAGES} images (4–6 recommended)")
    _check_refs(db, body)
    p = Product(**body.model_dump(exclude={"slug", "variants"}), slug=unique_slug(db, Product, body.slug or body.name))
    db.add(p)
    db.flush()
    for vin in body.variants:
        if _sku_taken(db, vin.sku, None):
            raise HTTPException(409, f"SKU {vin.sku} already exists")
        v = ProductVariant(
            product_id=p.id, sku=vin.sku.strip().upper(), size=vin.size.strip(), colour=vin.colour.strip(),
            stock_qty=0, price_override=vin.price_override, low_stock_threshold=vin.low_stock_threshold, is_active=vin.is_active,
        )
        db.add(v)
        db.flush()
        if vin.stock_qty:
            change_stock(db, v.id, vin.stock_qty, "purchase", user_id=user.id, note="Opening stock")
    # images are stored last, after every check has passed; a failure removes the files already stored
    saved: list[str] = []
    try:
        for n, f in enumerate(files):
            url = await save_image(f, "products")
            saved.append(url)
            db.add(ProductImage(product_id=p.id, url=url, sort_order=n))
        audit(db, user, "create", "product", p.id, {"name": p.name, "images": len(saved)}, request)
        db.commit()
    except Exception:
        db.rollback()
        for url in saved:
            delete_media(url)
        raise
    return product_detail(_load(db, p.id), admin=True)


@router.put("/products/{pid}")
def update_product(pid: int, body: ProductIn, request: Request, db: Session = Depends(get_db), user: User = Depends(staff_or_admin)):
    p = _load(db, pid)
    _check_refs(db, body)
    if body.is_active and not p.images:
        raise HTTPException(400, "Add at least one product image before making the product active")
    for k, v in body.model_dump(exclude={"slug", "variants"}).items():
        setattr(p, k, v)
    p.slug = unique_slug(db, Product, body.slug or body.name, exclude_id=p.id)
    existing = {v.id: v for v in p.variants}
    keep_ids = set()
    for vin in body.variants:
        if _sku_taken(db, vin.sku, vin.id):
            raise HTTPException(409, f"SKU {vin.sku} already exists")
        if vin.id and vin.id in existing:
            v = existing[vin.id]
            v.sku, v.size, v.colour = vin.sku.strip().upper(), vin.size.strip(), vin.colour.strip()
            v.price_override, v.low_stock_threshold, v.is_active = vin.price_override, vin.low_stock_threshold, vin.is_active
            keep_ids.add(v.id)
        else:
            v = ProductVariant(
                product_id=p.id, sku=vin.sku.strip().upper(), size=vin.size.strip(), colour=vin.colour.strip(), stock_qty=0,
                price_override=vin.price_override, low_stock_threshold=vin.low_stock_threshold, is_active=vin.is_active,
            )
            db.add(v)
            db.flush()
            keep_ids.add(v.id)
            if vin.stock_qty:
                change_stock(db, v.id, vin.stock_qty, "purchase", user_id=user.id, note="Opening stock")
    for vid, v in existing.items():
        if vid not in keep_ids:
            used = db.scalar(select(OrderItem.id).where(OrderItem.variant_id == vid).limit(1))
            if used or v.stock_qty:
                v.is_active = False  # keep history; hide from store
            else:
                db.delete(v)
    audit(db, user, "update", "product", p.id, {"name": p.name}, request)
    db.commit()
    db.expire_all()
    return product_detail(_load(db, p.id), admin=True)


@router.delete("/products/{pid}")
def delete_product(pid: int, request: Request, db: Session = Depends(get_db), user: User = Depends(staff_or_admin)):
    p = _load(db, pid)
    if db.scalar(select(OrderItem.id).where(OrderItem.product_id == pid).limit(1)):
        p.is_active = False
        audit(db, user, "archive", "product", pid, {"name": p.name}, request)
        db.commit()
        return {"ok": True, "archived": True, "message": "Product has orders, so it was moved to draft instead of deleted"}
    for img in p.images:
        delete_media(img.url)
    db.delete(p)
    audit(db, user, "delete", "product", pid, {"name": p.name}, request)
    db.commit()
    return {"ok": True}


@router.post("/products/{pid}/images")
async def upload_images(pid: int, files: list[UploadFile] = File(...), db: Session = Depends(get_db)):
    p = _load(db, pid)
    files = [f for f in files if f.filename]
    if len(p.images) + len(files) > MAX_IMAGES:
        raise HTTPException(400, f"A product can have at most {MAX_IMAGES} images (4–6 recommended)")
    start = max((i.sort_order for i in p.images), default=-1) + 1
    # a bad file part-way through must not leave the earlier files stored with nothing pointing at them
    saved: list[str] = []
    try:
        for n, f in enumerate(files):
            url = await save_image(f, "products")
            saved.append(url)
            db.add(ProductImage(product_id=p.id, url=url, sort_order=start + n))
        db.commit()
    except Exception:
        db.rollback()
        for url in saved:
            delete_media(url)
        raise
    db.expire_all()
    return product_detail(_load(db, pid), admin=True)


@router.put("/products/{pid}/images/order")
def reorder_images(pid: int, body: ImageOrderIn, db: Session = Depends(get_db)):
    p = _load(db, pid)
    pos = {iid: i for i, iid in enumerate(body.ids)}
    for img in p.images:
        if img.id in pos:
            img.sort_order = pos[img.id]
    db.commit()
    db.expire_all()
    return product_detail(_load(db, pid), admin=True)


@router.delete("/products/{pid}/images/{image_id}")
def delete_image(pid: int, image_id: int, db: Session = Depends(get_db)):
    # lock the product row so two deletes at the same moment cannot both pass the last-image check
    if not db.scalar(select(Product.id).where(Product.id == pid).with_for_update()):
        raise HTTPException(404, "Product not found")
    img = db.get(ProductImage, image_id)
    if not img or img.product_id != pid:
        raise HTTPException(404, "Image not found")
    if db.scalar(select(func.count(ProductImage.id)).where(ProductImage.product_id == pid)) <= 1:
        raise HTTPException(400, "A product must have at least one image. Add another image before deleting this one.")
    delete_media(img.url)
    db.delete(img)
    db.commit()
    db.expire_all()
    return product_detail(_load(db, pid), admin=True)


# ------------------------------------------------------------------ bulk import / export (AD-05)

COLUMNS = [
    "product_slug", "name", "brand", "category", "gender", "description", "material", "care", "hsn_code", "gst_rate",
    "mrp", "price", "tags", "is_active", "sku", "size", "colour", "stock_qty", "price_override", "low_stock_threshold",
]


def _export_rows(db: Session) -> list[list]:
    rows = []
    products = db.scalars(
        select(Product).options(selectinload(Product.variants), selectinload(Product.brand), selectinload(Product.category)).order_by(Product.id)
    ).all()
    for p in products:
        base = [p.slug, p.name, p.brand.name if p.brand else "", p.category.name if p.category else "", p.gender or "",
                p.description or "", p.material or "", p.care or "", p.hsn_code or "", p.gst_rate, p.mrp, p.price,
                p.tags or "", "yes" if p.is_active else "no"]
        if not p.variants:
            rows.append(base + ["", "", "", "", "", ""])
        for v in p.variants:
            rows.append(base + [v.sku, v.size, v.colour, v.stock_qty, v.price_override if v.price_override is not None else "", v.low_stock_threshold])
    return rows


@router.get("/products-export")
def export_products(format: str = Query(default="csv", pattern="^(csv|xlsx)$"), db: Session = Depends(get_db)):
    rows = _export_rows(db)
    if format == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(COLUMNS)
        w.writerows(rows)
        data = io.BytesIO(buf.getvalue().encode("utf-8-sig"))
        return StreamingResponse(data, media_type="text/csv", headers={"Content-Disposition": 'attachment; filename="products.csv"'})
    wb = Workbook()
    ws = wb.active
    ws.title = "products"
    ws.append(COLUMNS)
    for r in rows:
        ws.append([float(x) if isinstance(x, Decimal) else x for x in r])
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="products.xlsx"'})


def _read_upload(file: UploadFile, content: bytes) -> list[dict]:
    name = (file.filename or "").lower()
    if name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it)]
        return [dict(zip(header, ["" if c is None else c for c in row])) for row in it if any(c not in (None, "") for c in row)]
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig")
        return [r for r in csv.DictReader(io.StringIO(text)) if any((v or "").strip() for v in r.values())]
    raise HTTPException(400, "Upload a .csv or .xlsx file")


def _dec(v, field: str) -> Decimal | None:
    if v in (None, ""):
        return None
    try:
        return Decimal(str(v).strip())
    except InvalidOperation as exc:
        raise ValueError(f"{field} must be a number") from exc


@router.post("/products-import")
async def import_products(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db), user: User = Depends(staff_or_admin)):
    """Upsert by product_slug and sku. stock_qty sets the absolute stock (logged as a correction)."""
    rows = _read_upload(file, await file.read())
    created = updated = variants = 0
    errors: list[str] = []
    warnings: list[str] = []
    drafted: set[int] = set()
    brands = {b.name.lower(): b for b in db.scalars(select(Brand)).all()}
    cats = {c.name.lower(): c for c in db.scalars(select(Category)).all()}
    for n, row in enumerate(rows, start=2):
        r = {k.strip(): (str(v).strip() if v is not None else "") for k, v in row.items() if k}
        draft_note = None
        try:
            with db.begin_nested():
                if not r.get("name"):
                    raise ValueError("name is required")
                mrp, price = _dec(r.get("mrp"), "mrp"), _dec(r.get("price"), "price")
                if mrp is None or price is None:
                    raise ValueError("mrp and price are required")
                if price > mrp:
                    raise ValueError("price cannot be more than mrp")
                slug = r.get("product_slug") or None
                p = db.scalar(select(Product).where(Product.slug == slug)) if slug else None
                brand = None
                if r.get("brand"):
                    brand = brands.get(r["brand"].lower())
                    if not brand:
                        brand = Brand(name=r["brand"], slug=unique_slug(db, Brand, r["brand"]))
                        db.add(brand)
                        db.flush()
                        brands[r["brand"].lower()] = brand
                # same mandatory fields as the product form (ProductIn)
                if mrp <= 0 or price <= 0:
                    raise ValueError("mrp and price must be more than 0")
                if not r.get("category"):
                    raise ValueError("category is required")
                cat = cats.get(r["category"].lower())
                if not cat:
                    raise ValueError(f"category '{r['category']}' does not exist")
                gender = r.get("gender", "").lower()
                if gender not in GENDERS:
                    raise ValueError("gender is required: men, women, kids or unisex")
                hsn = r.get("hsn_code", "")
                if not re.fullmatch(r"\d{4,8}", hsn):
                    raise ValueError("hsn_code is required (4 to 8 digits)")
                gst_rate = _dec(r.get("gst_rate"), "gst_rate")
                if gst_rate is None:
                    raise ValueError("gst_rate is required")
                check_gst_rate(gst_rate)
                if not p and not r.get("sku"):
                    raise ValueError("sku, size and colour are required (every product needs at least one variant)")
                fields = dict(
                    name=r["name"], brand_id=brand.id if brand else None, category_id=cat.id,
                    gender=gender, description=r.get("description") or None, material=r.get("material") or None,
                    care=r.get("care") or None, hsn_code=hsn, gst_rate=gst_rate, mrp=mrp, price=price,
                    tags=r.get("tags") or None, is_active=r.get("is_active", "").lower() in ("yes", "true", "1", "active"),
                )
                if p:
                    for k, v in fields.items():
                        setattr(p, k, v)
                    updated += 1
                else:
                    p = Product(**fields, slug=unique_slug(db, Product, slug or r["name"]))
                    db.add(p)
                    db.flush()
                    created += 1
                # an image is mandatory and a file import cannot carry one, so a product without images stays a draft
                if p.is_active and not p.images:
                    p.is_active = False
                    draft_note = (p.id, f"Row {n}: '{p.name}' was saved as a draft. Add at least one image in Admin → Products to make it active.")
                if r.get("sku"):
                    if not r.get("size") or not r.get("colour"):
                        raise ValueError("size and colour are required with sku")
                    sku = r["sku"].upper()
                    v = db.scalar(select(ProductVariant).where(ProductVariant.sku == sku))
                    if v and v.product_id != p.id:
                        raise ValueError(f"SKU {sku} belongs to another product")
                    if not v:
                        v = ProductVariant(product_id=p.id, sku=sku, size=r["size"], colour=r["colour"], stock_qty=0)
                        db.add(v)
                        db.flush()
                    v.size, v.colour = r["size"], r["colour"]
                    v.price_override = _dec(r.get("price_override"), "price_override")
                    if v.price_override is not None and not 0 < v.price_override <= mrp:
                        raise ValueError("price_override must be more than 0 and not more than mrp")
                    if r.get("low_stock_threshold"):
                        v.low_stock_threshold = int(float(r["low_stock_threshold"]))
                    if r.get("stock_qty") != "":
                        target = int(float(r["stock_qty"]))
                        if target < 0:
                            raise ValueError("stock_qty cannot be negative")
                        if target != v.stock_qty:
                            change_stock(db, v.id, target - v.stock_qty, "correction", user_id=user.id, note="Bulk import")
                    variants += 1
            # noted only once the row has been saved
            if draft_note and draft_note[0] not in drafted:
                drafted.add(draft_note[0])
                warnings.append(draft_note[1])
        except (ValueError, HTTPException) as exc:
            errors.append(f"Row {n}: {getattr(exc, 'detail', exc)}")
    audit(db, user, "import", "product", None, {"created": created, "updated": updated, "variants": variants, "errors": len(errors)}, request)
    db.commit()
    return {"created": created, "updated": updated, "variants": variants, "errors": errors[:200], "warnings": warnings[:200]}
